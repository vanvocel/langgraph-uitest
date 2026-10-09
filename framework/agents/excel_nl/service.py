"""Read requirement Excel and write runs/<req>/nl/*.nl.md."""

from __future__ import annotations

import re
from dataclasses import asdict
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from framework.agents.excel_nl.config import (
    ExcelToNlConfig,
    header_aliases,
    load_excel_to_nl_config,
    write_req_excel_nl_template,
)
from framework.paths import init_requirement, req_dir
from framework.runner.excel_backfill import resolve_source_excel
from framework.runner.excel_backfill import load_excel_backfill_config

SkipKind = Literal["ai_outbound", "needs_fixture", ""]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_columns(ws: Worksheet, header_row: int) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, (ws.max_column or 1) + 1):
        val = ws.cell(header_row, col).value
        if val is None:
            continue
        headers[str(val).strip()] = col

    lower = {k.lower(): v for k, v in headers.items()}
    aliases = header_aliases()
    mapping: dict[str, int] = {}
    for key, names in aliases.items():
        for name in names:
            if name in headers:
                mapping[key] = headers[name]
                break
            if name.lower() in lower:
                mapping[key] = lower[name.lower()]
                break
    required = ["case_id", "title", "steps", "expect"]
    missing = [k for k in required if k not in mapping]
    if missing:
        raise ValueError(
            f"Excel missing columns {missing}; found headers={list(headers.keys())}. "
            "Need at least 用例ID/用例标题/测试步骤/预期结果 (MG0557-compatible)."
        )
    return mapping


def _detect_skip(
    *,
    case_id: str,
    title: str,
    pre: str,
    steps: str,
    expect: str,
    note: str,
    cfg: ExcelToNlConfig,
) -> SkipKind:
    if case_id in set(cfg.ai_outbound_ids or []):
        return "ai_outbound"
    if case_id in set(cfg.needs_fixture_ids or []):
        return "needs_fixture"
    blob = "\n".join([title, pre, steps, expect, note])
    if "【人工/AI外呼】" in blob or "【ai_outbound】" in blob.lower():
        return "ai_outbound"
    if "【需造数】" in blob or "【needs_fixture】" in blob:
        return "needs_fixture"
    # Heuristic keywords (title/note/pre stronger than steps)
    focus = f"{title}\n{note}\n{pre}"
    if re.search(r"AI\s*外呼|人工/AI外呼|第三方.*回传|外呼回传", focus, re.I):
        return "ai_outbound"
    if re.search(r"需造数|故障注入|只读账号|Mock|跨模块|人工造数", focus, re.I):
        return "needs_fixture"
    return ""


def render_nl(
    *,
    case_id: str,
    title: str,
    pre: str,
    steps: str,
    expect: str,
    cfg: ExcelToNlConfig,
    skip: SkipKind = "",
) -> str:
    lines = [
        f"用例 {case_id}：{title}",
        f"已登录（账号 {cfg.account_ref}）。",
    ]
    if skip == "ai_outbound":
        lines.append(
            "【人工/AI外呼】依赖第三方 AI 外呼回传，自动集不执行；"
            "编译时 tags 含 manual,ai_outbound，meta.manual_only=true。"
        )
    elif skip == "needs_fixture":
        lines.append(
            "【需造数】需造数/跨模块/故障注入/只读账号等前置，自动集暂不执行；"
            "编译时 tags 含 needs_fixture，meta.needs_fixture=true。"
        )
    page = (cfg.page_url or cfg.default_page_url or "").strip()
    tab = (cfg.tab or cfg.default_tab or "").strip()
    if page and tab:
        lines.append(f"页面入口：{page} ，进入页签【{tab}】。")
    elif page:
        lines.append(f"页面入口：{page} 。")
    elif tab:
        lines.append(f"进入页签【{tab}】。")
    if pre:
        lines.append(f"前置：{pre}")
    lines.append("步骤（请逐步编译为 YAML，勿省略边界值/校验点）：")
    lines.append(steps or "（Excel 未填写测试步骤）")
    lines.append(f"预期：{expect or '（Excel 未填写预期结果）'}")
    return "\n".join(lines) + "\n"


def excel_to_nl_requirement(
    req_id: str,
    *,
    excel: str | Path | None = None,
    page_url: str | None = None,
    tab: str | None = None,
    write_template: bool = True,
) -> dict[str, Any]:
    """Convert runs/<req>/excel → nl/*.nl.md. Returns a summary dict."""
    init_requirement(req_id)
    base = req_dir(req_id)
    if write_template:
        write_req_excel_nl_template(req_id)

    overrides: dict[str, Any] = {}
    if page_url is not None:
        overrides["page_url"] = page_url
    if tab is not None:
        overrides["tab"] = tab
    cfg = load_excel_to_nl_config(req_id, overrides=overrides or None)

    backfill_cfg = load_excel_backfill_config()
    source = resolve_source_excel(req_id, cfg=backfill_cfg, explicit=excel)
    if source is None:
        raise FileNotFoundError(
            f"no Excel found for {req_id}; put cases under runs/{req_id}/excel/cases.xlsx"
        )

    # data_only + non-read_only so merged/header detection stays simple and reliable
    wb = load_workbook(source, read_only=False, data_only=True)
    ws = wb[cfg.sheet] if cfg.sheet and cfg.sheet in wb.sheetnames else wb.active
    assert ws is not None
    cols = _find_columns(ws, cfg.header_row)

    nl_dir = base / "nl"
    nl_dir.mkdir(parents=True, exist_ok=True)
    written: list[str] = []
    skipped: dict[str, list[str]] = {"ai_outbound": [], "needs_fixture": []}
    blank = 0
    max_row = ws.max_row or cfg.data_start_row

    for row in range(cfg.data_start_row, max_row + 1):
        case_id = _cell_text(ws.cell(row, cols["case_id"]).value)
        if not case_id:
            blank += 1
            continue
        title = _cell_text(ws.cell(row, cols["title"]).value) if "title" in cols else case_id
        pre = _cell_text(ws.cell(row, cols["precondition"]).value) if "precondition" in cols else ""
        steps = _cell_text(ws.cell(row, cols["steps"]).value)
        expect = _cell_text(ws.cell(row, cols["expect"]).value)
        note = _cell_text(ws.cell(row, cols["note"]).value) if "note" in cols else ""
        skip = _detect_skip(
            case_id=case_id,
            title=title,
            pre=pre,
            steps=steps,
            expect=expect,
            note=note,
            cfg=cfg,
        )
        if skip:
            skipped[skip].append(case_id)
        text = render_nl(
            case_id=case_id,
            title=title,
            pre=pre,
            steps=steps,
            expect=expect,
            cfg=cfg,
            skip=skip,
        )
        out = nl_dir / f"{case_id}.nl.md"
        out.write_text(text, encoding="utf-8")
        written.append(case_id)

    wb.close()
    return {
        "ok": True,
        "req_id": req_id,
        "excel": str(source),
        "nl_dir": str(nl_dir),
        "written": written,
        "count": len(written),
        "blank_rows": blank,
        "skipped": skipped,
        "page_url": cfg.page_url or cfg.default_page_url,
        "tab": cfg.tab or cfg.default_tab,
        "config": {k: v for k, v in asdict(cfg).items() if k not in {"ai_outbound_ids", "needs_fixture_ids"}},
    }
