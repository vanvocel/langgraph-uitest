"""Backfill requirement Excel with pass / fail / Block + screenshots after a run."""

from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from framework.paths import load_settings, req_dir
from framework.schema.loader import load_case

RESULT_PASS = "通过"
RESULT_FAIL = "失败"
RESULT_BLOCK = "Block"

_FILL = {
    RESULT_PASS: PatternFill("solid", fgColor="C6EFCE"),
    RESULT_FAIL: PatternFill("solid", fgColor="FFC7CE"),
    RESULT_BLOCK: PatternFill("solid", fgColor="FFEB9C"),
}
_FONT = {
    RESULT_PASS: Font(color="006100", bold=True),
    RESULT_FAIL: Font(color="9C0006", bold=True),
    RESULT_BLOCK: Font(color="9C5700", bold=True),
}

_HEADER_ALIASES = {
    "case_id": ("用例ID", "用例编号", "case_id", "Case ID", "用例id"),
    "result": ("测试结果", "执行结果", "结果", "result", "status"),
    "screenshot": ("测试截图", "截图", "screenshot", "附件"),
    "note": ("测试说明", "备注", "说明", "note", "comment"),
}


@dataclass
class ExcelBackfillConfig:
    enabled: bool = True
    source: str | None = None
    output_name: str = "{req_id}_结果回填.xlsx"
    sheet: str | None = None
    header_row: int = 2
    data_start_row: int = 3
    case_id_header: str = "用例ID"
    result_header: str = "测试结果"
    screenshot_header: str = "测试截图"
    note_header: str = "测试说明"
    search_downloads: bool = True
    image_width: int = 180
    image_height: int = 100


def load_excel_backfill_config(overrides: dict[str, Any] | None = None) -> ExcelBackfillConfig:
    raw = (load_settings().get("excel_backfill") or {}).copy()
    if overrides:
        raw.update({k: v for k, v in overrides.items() if v is not None})
    known = {f.name for f in ExcelBackfillConfig.__dataclass_fields__.values()}  # type: ignore[attr-defined]
    return ExcelBackfillConfig(**{k: v for k, v in raw.items() if k in known})


def _normalize_case_id(raw: Any) -> str:
    if raw is None:
        return ""
    s = str(raw).strip().strip("'\"")
    if not s:
        return ""
    # pytest may stringify Paths
    if "WindowsPath(" in s or "PosixPath(" in s:
        m = re.search(r"['\"]([^'\"]+)['\"]", s)
        if m:
            s = m.group(1)
    if "/" in s or "\\" in s:
        s = Path(s).stem
    # strip surrounding quotes again after Path
    return s.strip().strip("'\"")


def _find_header_columns(ws: Worksheet, header_row: int, cfg: ExcelBackfillConfig) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, (ws.max_column or 1) + 1):
        val = ws.cell(header_row, col).value
        if val is None:
            continue
        text = str(val).strip()
        headers[text] = col

    def pick(key: str, preferred: str) -> int | None:
        if preferred in headers:
            return headers[preferred]
        for alias in _HEADER_ALIASES[key]:
            if alias in headers:
                return headers[alias]
        # case-insensitive
        lower = {k.lower(): v for k, v in headers.items()}
        if preferred.lower() in lower:
            return lower[preferred.lower()]
        for alias in _HEADER_ALIASES[key]:
            if alias.lower() in lower:
                return lower[alias.lower()]
        return None

    mapping = {
        "case_id": pick("case_id", cfg.case_id_header),
        "result": pick("result", cfg.result_header),
        "screenshot": pick("screenshot", cfg.screenshot_header),
        "note": pick("note", cfg.note_header),
    }
    missing = [k for k, v in mapping.items() if v is None]
    if missing:
        raise ValueError(
            f"Excel missing required columns {missing}; headers={list(headers.keys())}"
        )
    return {k: int(v) for k, v in mapping.items()}  # type: ignore[arg-type]


def resolve_source_excel(
    req_id: str,
    *,
    cfg: ExcelBackfillConfig,
    explicit: str | Path | None = None,
) -> Path | None:
    if explicit:
        path = Path(explicit)
        return path if path.is_file() else None

    base = req_dir(req_id)
    if cfg.source:
        candidate = Path(cfg.source.format(req_id=req_id))
        if not candidate.is_absolute():
            candidate = base / candidate
        if candidate.is_file():
            return candidate

    excel_dir = base / "excel"
    if excel_dir.is_dir():
        preferred = excel_dir / "cases.xlsx"
        if preferred.is_file():
            return preferred
        files = sorted(
            p
            for p in excel_dir.glob("*.xlsx")
            if p.is_file() and "结果回填" not in p.name and not p.name.startswith("~$")
        )
        if files:
            return files[0]

    if cfg.search_downloads:
        downloads = Path.home() / "Downloads"
        if downloads.is_dir():
            hits = sorted(
                p
                for p in downloads.glob(f"{req_id}*.xlsx")
                if p.is_file() and "结果回填" not in p.name and not p.name.startswith("~$")
            )
            if hits:
                return hits[0]
    return None


def load_allure_by_case(allure_dir: Path) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    if not allure_dir.is_dir():
        return out
    for path in allure_dir.glob("*-result.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        params = {p["name"]: p.get("value") for p in data.get("parameters") or []}
        case_id = _normalize_case_id(params.get("case_id"))
        if not case_id:
            case_id = _normalize_case_id(params.get("yaml_case_path"))
        if not case_id:
            continue
        imgs: list[str] = []
        for step in data.get("steps") or []:
            for att in step.get("attachments") or []:
                if str(att.get("type") or "").startswith("image") and att.get("source"):
                    imgs.append(str(att["source"]))
        for att in data.get("attachments") or []:
            if str(att.get("type") or "").startswith("image") and att.get("source"):
                imgs.append(str(att["source"]))
        out[case_id] = {
            "status": data.get("status") or "unknown",
            "shot": imgs[-1] if imgs else None,
            "name": data.get("name") or "",
            "error": ((data.get("statusDetails") or {}).get("message") or ""),
        }
    return out


def classify_yaml_case(yaml_dir: Path, case_id: str) -> str:
    """Return ui / manual / needs_fixture / missing."""
    path = yaml_dir / f"{case_id}.yaml"
    if not path.is_file():
        # also accept .yml
        path = yaml_dir / f"{case_id}.yml"
    if not path.is_file():
        return "missing"
    try:
        case = load_case(path)
    except Exception:  # noqa: BLE001
        text = path.read_text(encoding="utf-8")
        if "ai_outbound" in text or "manual_only" in text:
            return "manual"
        if "needs_fixture" in text:
            return "needs_fixture"
        return "ui"
    tags = {str(t).lower() for t in (case.tags or [])}
    meta = case.meta or {}
    if tags & {"manual", "ai_outbound"} or meta.get("manual_only"):
        return "manual"
    if "needs_fixture" in tags or meta.get("needs_fixture"):
        return "needs_fixture"
    return "ui"


def resolve_case_result(
    case_id: str,
    *,
    yaml_dir: Path,
    allure_map: dict[str, dict[str, Any]],
    allure_dir: Path,
    shot_dir: Path,
) -> tuple[str, str, Path | None]:
    kind = classify_yaml_case(yaml_dir, case_id)
    if kind == "manual":
        return RESULT_BLOCK, "人工/AI外呼等用例，等待人工执行（自动集已跳过）", None
    if kind == "needs_fixture":
        return RESULT_BLOCK, "需造数/跨模块/故障注入等前置，暂不自动执行", None
    if kind == "missing":
        return RESULT_BLOCK, "Excel 用例无对应 YAML", None

    info = allure_map.get(case_id)
    if not info:
        return RESULT_BLOCK, "本次未执行（无 Allure 结果）", None

    shot_path: Path | None = None
    if info.get("shot"):
        src = allure_dir / str(info["shot"])
        if src.is_file():
            shot_dir.mkdir(parents=True, exist_ok=True)
            dest = shot_dir / f"{case_id}.png"
            shutil.copy2(src, dest)
            shot_path = dest

    status = info["status"]
    if status == "passed":
        return RESULT_PASS, "UI自动化执行通过", shot_path
    if status in {"failed", "broken"}:
        note = (info.get("error") or "执行失败").strip().splitlines()[0][:200]
        return RESULT_FAIL, note or "执行失败", shot_path
    if status == "skipped":
        return RESULT_BLOCK, "执行被跳过", shot_path
    return RESULT_BLOCK, f"未知状态: {status}", shot_path


def backfill_excel(
    *,
    source: Path,
    output: Path,
    yaml_dir: Path,
    allure_dir: Path,
    cfg: ExcelBackfillConfig | None = None,
) -> dict[str, Any]:
    """Copy source workbook, write results/screenshots/notes, save to output."""
    cfg = cfg or ExcelBackfillConfig()
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)

    wb = load_workbook(output)
    ws = wb[cfg.sheet] if cfg.sheet and cfg.sheet in wb.sheetnames else wb.active
    cols = _find_header_columns(ws, cfg.header_row, cfg)

    allure_map = load_allure_by_case(allure_dir)
    shot_dir = output.parent / "excel_screenshots"
    stats = {RESULT_PASS: 0, RESULT_FAIL: 0, RESULT_BLOCK: 0}
    rows_written = 0

    ws.column_dimensions[get_column_letter(cols["result"])].width = 12
    ws.column_dimensions[get_column_letter(cols["screenshot"])].width = 28
    ws.column_dimensions[get_column_letter(cols["note"])].width = 48

    for row in range(cfg.data_start_row, (ws.max_row or cfg.data_start_row) + 1):
        case_raw = ws.cell(row, cols["case_id"]).value
        case_id = _normalize_case_id(case_raw)
        if not case_id:
            continue
        result, note, shot = resolve_case_result(
            case_id,
            yaml_dir=yaml_dir,
            allure_map=allure_map,
            allure_dir=allure_dir,
            shot_dir=shot_dir,
        )
        stats[result] = stats.get(result, 0) + 1
        rows_written += 1

        cell_r = ws.cell(row, cols["result"], result)
        cell_r.fill = _FILL.get(result, PatternFill())
        cell_r.font = _FONT.get(result, Font())
        cell_r.alignment = Alignment(horizontal="center", vertical="center")

        existing = ws.cell(row, cols["note"]).value
        if existing and str(existing).strip() and not str(existing).startswith(note):
            ws.cell(row, cols["note"], f"{note}｜原说明: {existing}")
        else:
            ws.cell(row, cols["note"], note)
        ws.cell(row, cols["note"]).alignment = Alignment(wrap_text=True, vertical="center")

        ws.cell(row, cols["screenshot"]).value = None
        if shot and shot.is_file():
            img = XLImage(str(shot))
            img.width = cfg.image_width
            img.height = cfg.image_height
            ws.row_dimensions[row].height = max(80, cfg.image_height * 0.75)
            img.anchor = f"{get_column_letter(cols['screenshot'])}{row}"
            ws.add_image(img)
        else:
            ws.row_dimensions[row].height = 30

    wb.save(output)
    summary = {
        "ok": True,
        "excel": str(output),
        "source": str(source),
        "stats": stats,
        "rows": rows_written,
        "allure_cases": len(allure_map),
    }
    (output.parent / "excel_backfill_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return summary


def backfill_requirement_excel(
    req_id: str,
    *,
    source: str | Path | None = None,
    allure_dir: Path | None = None,
    enabled: bool | None = None,
    config_overrides: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """High-level entry used by CLI after pytest run."""
    cfg = load_excel_backfill_config(config_overrides)
    if enabled is not None:
        cfg.enabled = enabled
    if not cfg.enabled:
        return {"ok": False, "skipped": True, "reason": "excel_backfill disabled"}

    base = req_dir(req_id)
    src = resolve_source_excel(req_id, cfg=cfg, explicit=source)
    if src is None:
        return {
            "ok": False,
            "skipped": True,
            "reason": (
                f"no Excel source for {req_id}; put file under runs/{req_id}/excel/ "
                f"or Downloads/{req_id}*.xlsx"
            ),
        }

    out_name = cfg.output_name.format(req_id=req_id)
    output = base / out_name
    results = allure_dir or (base / "allure-results")
    return backfill_excel(
        source=src,
        output=output,
        yaml_dir=base / "yaml",
        allure_dir=results,
        cfg=cfg,
    )
