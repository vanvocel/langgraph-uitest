"""Unit tests for Excel result backfill."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from framework.runner.excel_backfill import (
    RESULT_BLOCK,
    RESULT_FAIL,
    RESULT_PASS,
    ExcelBackfillConfig,
    backfill_excel,
    classify_yaml_case,
    load_allure_by_case,
)


def _write_source(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "冒烟测试用例"
    ws.append(["标题"])
    ws.append(
        [
            "钉钉需求编号",
            "模块",
            "用例ID",
            "用例标题",
            "前置条件",
            "测试步骤",
            "预期结果",
            "测试结果",
            "测试截图",
            "测试说明",
        ]
    )
    ws.append(["R1", "m", "TC-DEMO-001", "ui pass", "", "", "", "", "", "orig"])
    ws.append(["R1", "m", "TC-DEMO-002", "ui fail", "", "", "", "", "", ""])
    ws.append(["R1", "m", "TC-DEMO-003", "manual", "", "", "", "", "", ""])
    wb.save(path)


def _write_yaml(yaml_dir: Path, case_id: str, tags: list[str], meta: dict | None = None) -> None:
    meta = meta or {}
    lines = [
        f"case_id: {case_id}",
        "title: demo",
        "requirement_id: DEMO",
        "base_url: https://example.com",
        "tags:",
        *[f"- {t}" for t in tags],
        "meta:",
        *[f"  {k}: {json.dumps(v) if isinstance(v, bool) else v}" for k, v in meta.items()],
        "preconditions: []",
        "steps:",
        "- action: open",
        "  url: https://example.com",
    ]
    (yaml_dir / f"{case_id}.yaml").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_allure(allure_dir: Path, case_id: str, status: str, shot_name: str | None) -> None:
    atts = []
    if shot_name:
        png = allure_dir / shot_name
        png.write_bytes(
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
            b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f"
            b"\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
        )
        atts = [{"name": "shot", "source": shot_name, "type": "image/png"}]
    payload = {
        "name": case_id,
        "status": status,
        "parameters": [{"name": "case_id", "value": f"'{case_id}'"}],
        "steps": [{"name": "assert", "attachments": atts}],
        "statusDetails": {"message": "boom"} if status != "passed" else {},
    }
    (allure_dir / f"{case_id}-result.json").write_text(
        json.dumps(payload, ensure_ascii=False), encoding="utf-8"
    )


def test_backfill_pass_fail_block(tmp_path: Path):
    source = tmp_path / "cases.xlsx"
    out = tmp_path / "out.xlsx"
    yaml_dir = tmp_path / "yaml"
    allure_dir = tmp_path / "allure-results"
    yaml_dir.mkdir()
    allure_dir.mkdir()
    _write_source(source)
    _write_yaml(yaml_dir, "TC-DEMO-001", ["ui"])
    _write_yaml(yaml_dir, "TC-DEMO-002", ["ui"])
    _write_yaml(yaml_dir, "TC-DEMO-003", ["manual", "ai_outbound"], {"manual_only": True})
    _write_allure(allure_dir, "TC-DEMO-001", "passed", "a.png")
    _write_allure(allure_dir, "TC-DEMO-002", "failed", "b.png")

    summary = backfill_excel(
        source=source,
        output=out,
        yaml_dir=yaml_dir,
        allure_dir=allure_dir,
        cfg=ExcelBackfillConfig(),
    )
    assert summary["ok"] is True
    assert summary["stats"][RESULT_PASS] == 1
    assert summary["stats"][RESULT_FAIL] == 1
    assert summary["stats"][RESULT_BLOCK] == 1

    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(3, 8).value == RESULT_PASS
    assert ws.cell(4, 8).value == RESULT_FAIL
    assert ws.cell(5, 8).value == RESULT_BLOCK
    assert "原说明" in str(ws.cell(3, 10).value)
    assert len(ws._images) == 2


def test_classify_and_allure_loader(tmp_path: Path):
    yaml_dir = tmp_path / "yaml"
    yaml_dir.mkdir()
    _write_yaml(yaml_dir, "A", ["needs_fixture"], {"needs_fixture": True})
    assert classify_yaml_case(yaml_dir, "A") == "needs_fixture"
    assert classify_yaml_case(yaml_dir, "missing") == "missing"

    allure_dir = tmp_path / "ar"
    allure_dir.mkdir()
    _write_allure(allure_dir, "TC-X-1", "passed", None)
    m = load_allure_by_case(allure_dir)
    assert "TC-X-1" in m
    assert m["TC-X-1"]["status"] == "passed"
